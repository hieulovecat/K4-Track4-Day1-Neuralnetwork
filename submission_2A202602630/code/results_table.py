"""results_table.py — lưu kết quả từng lần chạy ra JSON, rồi điền experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx.

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu tự tính, không ghi đè)
"""
from __future__ import annotations

import json
import math
from pathlib import Path

import numpy as np
import openpyxl

DATA_COLUMNS = [
    "exp_id", "group", "description", "loss", "optimizer", "lr", "weight_decay", "batch", "epochs",
    "hidden", "dropout", "clip_norm", "precision", "init", "seed", "step0_loss", "best_val_loss",
    "best_epoch", "final_train_loss", "final_val_loss", "val_acc", "val_macro_f1", "time_per_epoch_s",
    "peak_mem_MB", "diverged", "eval_acc", "eval_macro_f1", "figure_file", "notes",
]
FORMULA_COLUMNS = ["step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise"]
GROUPS = ("baseline", "loss", "optimizer", "hparam", "dropout", "clipping", "amp", "init", "final", "other")

# Giá trị trong cfg -> giá trị đúng danh sách chọn (data validation) của mẫu
LOSS_NAMES = {"ce": "CE", "mse": "MSE"}
OPT_NAMES = {"sgd": "SGD", "sgd_momentum": "SGD+momentum", "adam": "Adam", "adamw": "AdamW"}

FIRST_ROW, LAST_ROW = 2, 61          # mẫu có công thức và danh sách chọn cho 60 dòng
SEED_ROWS = range(2, 7)              # sheet Seeds: A2:A6
SUMMARY_NOTE_COL = 8                 # sheet Summary: cột H "nhận xét ngắn"


def _clean(obj):
    """Đưa về kiểu JSON chuẩn: tuple -> list, numpy -> python, NaN/inf -> None."""
    if isinstance(obj, dict):
        return {k: _clean(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [_clean(v) for v in obj]
    if isinstance(obj, np.generic):
        obj = obj.item()
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    return obj


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi cfg, history, summary (KHÔNG ghi best_state) ra <results_dir>/<exp_id>.json. Trả về đường dẫn."""
    out = Path(results_dir)
    out.mkdir(parents=True, exist_ok=True)
    payload = _clean({k: result[k] for k in ("cfg", "history", "summary")})
    path = out / f"{result['cfg']['exp_id']}.json"
    with open(path, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=1)
    return str(path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    results = []
    for p in sorted(Path(results_dir).glob("*.json")):
        with open(p, encoding="utf-8") as f:
            r = json.load(f)
        r["cfg"]["hidden"] = tuple(r["cfg"]["hidden"])
        results.append(r)
    return sorted(results, key=lambda r: r["cfg"]["exp_id"])


def _num(v, ndigits: int = 6):
    """Số đo -> số làm tròn cho bảng; None/NaN -> ô trống."""
    if v is None or (isinstance(v, float) and not math.isfinite(v)):
        return None
    return round(float(v), ndigits)


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Một kết quả -> một dòng của bảng (khoá trùng tên cột).

    eval_scores: dict đọc từ eval_result.json (khoá "accuracy", "macro_f1"), CHỈ truyền cho
                 baseline và cấu hình cuối cùng.
    notes: ghép thêm sau cfg["notes"] (nếu có).
    """
    cfg, s = result["cfg"], result["summary"]
    if cfg.get("group") not in GROUPS:
        raise ValueError(f"{cfg['exp_id']}: group {cfg.get('group')!r} không thuộc {GROUPS}")

    extra = []  # những gì không có cột riêng nhưng ảnh hưởng tới kết quả -> ghi vào notes
    if cfg.get("scheduler"):
        extra.append(f"scheduler={cfg['scheduler']}" + (f", warmup={cfg['warmup_steps']} bước" if cfg.get("warmup_steps") else ""))
    if cfg.get("optimizer") == "sgd_momentum" and cfg.get("momentum", 0.9) != 0.9:
        extra.append(f"momentum={cfg['momentum']}")
    if s.get("clip_frac"):
        extra.append(f"clip kích hoạt ở {100 * s['clip_frac']:.1f}% số bước")
    if s.get("amp_skipped_steps"):
        extra.append(f"GradScaler bỏ qua {s['amp_skipped_steps']} bước")
    if s.get("diverged"):
        extra.append(f"dừng sớm sau {s.get('epochs_run')} epoch do loss NaN/inf")
    note_parts = [p for p in (cfg.get("notes", ""), notes, "; ".join(extra)) if p]

    return {
        "exp_id": cfg["exp_id"],
        "group": cfg["group"],
        "description": cfg.get("description", ""),
        "loss": LOSS_NAMES[cfg["loss"]],
        "optimizer": OPT_NAMES[cfg["optimizer"]],
        "lr": cfg["lr"],
        "weight_decay": cfg.get("weight_decay", 0),
        "batch": cfg["batch"],
        "epochs": cfg["epochs"],
        "hidden": "-".join(str(h) for h in cfg["hidden"]),
        "dropout": cfg.get("dropout", 0),
        "clip_norm": cfg["clip_norm"] if cfg.get("clip_norm") else "none",
        "precision": cfg["precision"],
        "init": cfg["init"],
        "seed": cfg["seed"],
        "step0_loss": _num(s.get("step0_loss")),
        "best_val_loss": _num(s.get("best_val_loss")),
        "best_epoch": s.get("best_epoch") or None,
        "final_train_loss": _num(s.get("final_train_loss")),
        "final_val_loss": _num(s.get("final_val_loss")),
        "val_acc": _num(s.get("val_acc")),
        "val_macro_f1": _num(s.get("val_macro_f1")),
        "time_per_epoch_s": _num(s.get("time_per_epoch_s"), 3),
        "peak_mem_MB": _num(s.get("peak_mem_MB"), 1),
        "diverged": "Y" if s.get("diverged") else "N",
        "eval_acc": _num(eval_scores["accuracy"]) if eval_scores else None,
        "eval_macro_f1": _num(eval_scores["macro_f1"]) if eval_scores else None,
        "figure_file": f"figures/{cfg['exp_id']}.png",
        "notes": " | ".join(note_parts),
    }


def write_xlsx(rows: list[dict], template_path: str, out_path: str,
               seed_ids: list[str] | None = None, summary_notes: dict[str, str] | None = None) -> None:
    """Điền các dòng vào sheet "Experiments" của mẫu (từ dòng 2), giữ nguyên công thức, định dạng và
    danh sách chọn, rồi lưu thành out_path.

    seed_ids      : exp_id các lần chạy baseline khác seed -> sheet Seeds, cột A (tối đa 5)
    summary_notes : {group: nhận xét} -> sheet Summary, cột "nhận xét ngắn"
    Sau khi lưu, mở file bằng Excel/LibreOffice một lần để các công thức tính lại.
    """
    n_slots = LAST_ROW - FIRST_ROW + 1
    if len(rows) > n_slots:
        raise ValueError(f"mẫu chỉ có công thức cho {n_slots} dòng, nhận {len(rows)}")
    ids = [r["exp_id"] for r in rows]
    if len(set(ids)) != len(ids):
        raise ValueError("exp_id bị trùng: " + ", ".join(sorted({i for i in ids if ids.count(i) > 1})))

    wb = openpyxl.load_workbook(template_path)  # không data_only: giữ công thức
    ws = wb["Experiments"]
    header = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    missing = [c for c in DATA_COLUMNS if c not in header]
    assert not missing, f"mẫu thiếu cột {missing}"

    # Xoá dòng baseline minh hoạ của mẫu (và mọi giá trị cũ) ở các cột dữ liệu; cột công thức giữ nguyên
    for r in range(FIRST_ROW, LAST_ROW + 1):
        for col in DATA_COLUMNS:
            ws.cell(r, header[col]).value = None
    for i, row in enumerate(rows):
        for col in DATA_COLUMNS:
            ws.cell(FIRST_ROW + i, header[col]).value = row.get(col)

    if seed_ids is not None:
        if len(seed_ids) > len(SEED_ROWS):
            raise ValueError(f"sheet Seeds chỉ có {len(SEED_ROWS)} ô")
        unknown = [s for s in seed_ids if s not in ids]
        if unknown:
            raise ValueError(f"seed_ids không có trong bảng: {unknown}")
        seeds = wb["Seeds"]
        for i, r in enumerate(SEED_ROWS):
            seeds.cell(r, 1).value = seed_ids[i] if i < len(seed_ids) else None

    if summary_notes:
        summ = wb["Summary"]
        group_row = {summ.cell(r, 1).value: r for r in range(2, summ.max_row + 1)}
        for g, text in summary_notes.items():
            if g not in group_row:
                raise ValueError(f"Summary không có nhóm {g!r}")
            summ.cell(group_row[g], SUMMARY_NOTE_COL).value = text

    # openpyxl không tính công thức; cờ này buộc Excel/LibreOffice tính lại toàn bộ khi mở file
    wb.calculation.fullCalcOnLoad = True
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
